import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import csv

# 초등학교 5학년 25명의 모의 관심사 데이터
# 스크린샷의 1~10번 색깔 데이터를 그대로 반영하고, 과목과 동물 등을 실감나게 구성
sample_students = [
    {"num": 1, "color": "빨강", "subject": "체육", "animal": "강아지", "food": "치킨", "hobby": "축구/운동"},
    {"num": 2, "color": "파랑", "subject": "과학", "animal": "고양이", "food": "피자", "hobby": "게임"},
    {"num": 3, "color": "초록", "subject": "미술", "animal": "토끼", "food": "떡볶이", "hobby": "그림그리기"},
    {"num": 4, "color": "노랑", "subject": "음악", "animal": "햄스터", "food": "마라탕", "hobby": "유튜브 시청"},
    {"num": 5, "color": "빨강", "subject": "체육", "animal": "호랑이", "food": "치킨", "hobby": "게임"},
    {"num": 6, "color": "노랑", "subject": "수학", "animal": "강아지", "food": "돈가스", "hobby": "만들기/블록"},
    {"num": 7, "color": "초록", "subject": "과학", "animal": "판다", "food": "라면", "hobby": "웹툰 보기"},
    {"num": 8, "color": "노랑", "subject": "사회", "animal": "고양이", "food": "떡볶이", "hobby": "친구와 놀기"},
    {"num": 9, "color": "초록", "subject": "미술", "animal": "돌고래", "food": "탕수육", "hobby": "음악 듣기"},
    {"num": 10, "color": "파랑", "subject": "실과", "animal": "강아지", "food": "치킨", "hobby": "컴퓨터/코딩"},
    {"num": 11, "color": "보라", "subject": "국어", "animal": "고양이", "food": "피자", "hobby": "독서/글쓰기"},
    {"num": 12, "color": "분홍", "subject": "음악", "animal": "토끼", "food": "마라탕", "hobby": "노래/춤"},
    {"num": 13, "color": "파랑", "subject": "체육", "animal": "치타", "food": "삼겹살", "hobby": "축구/운동"},
    {"num": 14, "color": "하늘", "subject": "영어", "animal": "돌고래", "food": "초밥", "hobby": "유튜브 시청"},
    {"num": 15, "color": "빨강", "subject": "체육", "animal": "사자", "food": "치킨", "hobby": "게임"},
    {"num": 16, "color": "주황", "subject": "미술", "animal": "여우", "food": "떡볶이", "hobby": "그림그리기"},
    {"num": 17, "color": "초록", "subject": "과학", "animal": "공룡", "food": "햄버거", "hobby": "레고 조립"},
    {"num": 18, "color": "노랑", "subject": "수학", "animal": "강아지", "food": "피자", "hobby": "보드게임"},
    {"num": 19, "color": "파랑", "subject": "실과", "animal": "고양이", "food": "치킨", "hobby": "로봇/코딩"},
    {"num": 20, "color": "보라", "subject": "음악", "animal": "펭귄", "food": "스파게티", "hobby": "악기 연주"},
    {"num": 21, "color": "빨강", "subject": "체육", "animal": "강아지", "food": "삼겹살", "hobby": "자전거 타기"},
    {"num": 22, "color": "초록", "subject": "사회", "animal": "수달", "food": "라면", "hobby": "유튜브 시청"},
    {"num": 23, "color": "파랑", "subject": "과학", "animal": "고래", "food": "돈가스", "hobby": "우주/과학탐구"},
    {"num": 24, "color": "분홍", "subject": "미술", "animal": "토끼", "food": "마라탕", "hobby": "다이어리 꾸미기"},
    {"num": 25, "color": "하양", "subject": "도덕", "animal": "백곰", "food": "짜장면", "hobby": "영화 보기"}
]

# 스타일 정의
header_font = Font(name="맑은 고딕", size=11, bold=True, color="1F2937")
data_font = Font(name="맑은 고딕", size=10, color="111827")
header_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
thin_border = Border(
    left=Side(style="thin", color="D1D5DB"),
    right=Side(style="thin", color="D1D5DB"),
    top=Side(style="thin", color="D1D5DB"),
    bottom=Side(style="thin", color="D1D5DB")
)
center_align = Alignment(horizontal="center", vertical="center")

import sys
import io
if sys.platform == "win32":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

def create_excel(filename, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "엔트리_데이터"
    
    # 헤더 작성
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        
    # 데이터 작성
    for r_idx, row_data in enumerate(rows, start=2):
        ws.append(row_data)
        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.font = data_font
            cell.alignment = center_align
            cell.border = thin_border

    # 열 너비 자동 조정
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len * 2.5, 14)
        
    wb.save(filename)
    print(f"[OK] 엑셀 파일 생성 완료: {filename}")

def create_csv(filename, headers, rows):
    # 엔트리 및 엑셀에서 한글이 깨지지 않도록 utf-8-sig(BOM) 인코딩 사용
    with open(filename, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        writer.writerows(rows)
    print(f"[OK] CSV 파일 생성 완료: {filename}")

# 1. 스크린샷과 정확히 동일한 형식: [친구번호, 색깔]
color_rows = [[s["num"], s["color"]] for s in sample_students]
create_excel("친구번호_색깔_데이터.xlsx", ["친구번호", "색깔"], color_rows)
create_csv("친구번호_색깔_데이터.csv", ["친구번호", "색깔"], color_rows)

# 2. 과목 데이터: [친구번호, 과목]
subject_rows = [[s["num"], s["subject"]] for s in sample_students]
create_excel("친구번호_과목_데이터.xlsx", ["친구번호", "과목"], subject_rows)
create_csv("친구번호_과목_데이터.csv", ["친구번호", "과목"], subject_rows)

# 3. 동물 데이터: [친구번호, 동물]
animal_rows = [[s["num"], s["animal"]] for s in sample_students]
create_excel("친구번호_동물_데이터.xlsx", ["친구번호", "동물"], animal_rows)
create_csv("친구번호_동물_데이터.csv", ["친구번호", "동물"], animal_rows)

# 4. 전체 종합 관심사 데이터: [친구번호, 색깔, 과목, 동물, 음식, 여가활동] (단일 전체 테이블)
total_headers = ["친구번호", "색깔", "과목", "동물", "좋아하는음식", "여가활동"]
total_rows = [[s["num"], s["color"], s["subject"], s["animal"], s["food"], s["hobby"]] for s in sample_students]

# 메인 전체 테이블 파일
create_excel("초등5학년_우리반_관심사_전체데이터.xlsx", total_headers, total_rows)
create_csv("초등5학년_우리반_관심사_전체데이터.csv", total_headers, total_rows)

# 호환용 파일
create_excel("초등5학년_관심사_종합데이터.xlsx", total_headers, total_rows)
create_csv("초등5학년_관심사_종합데이터.csv", total_headers, total_rows)

print("\n모든 엑셀 및 CSV 파일이 엔트리 규격에 맞춰 성공적으로 생성되었습니다!")


